"""BICS is a rotating-module survey, so a wave carries only the questions asked.

A wave workbook that omits a configured question sheet is normal -- the next
wave that asks it restores that question's full time series. Hard-failing on an
absent sheet took the collector offline whenever the latest wave happened to
rotate a question out (observed live: waves 162 and 164 carry one configured
question, wave 163 carries four). Genuine structural drift -- a workbook with
none of the configured questions -- must still fail loudly.
"""

from __future__ import annotations

from datetime import UTC, datetime

import pytest

from scripts import extract as E


def _workbook(sheet_names: list[str], rows: int = 0) -> bytes:
    """Build a minimal workbook shaped like a BICS release."""
    import io

    import openpyxl

    book = openpyxl.Workbook()
    index = book.active
    index.title = "Weighted Time Series"
    index["A1"] = "Released 15 September 2026"
    for name in sheet_names:
        sheet = book.create_sheet(name)
        sheet["A1"] = "Question: synthetic"
        responses = sorted(E.SHEETS[name])
        for offset, response in enumerate(responses):
            sheet.cell(6, 4 + offset, response)
        for column, label in enumerate(["Dates", "Wave", "Industry/Size Band"], start=1):
            sheet.cell(6, column, label)
        for index in range(rows):
            row = 7 + index
            sheet.cell(row, 1, "7 September 2026 to 20 September 2026")
            sheet.cell(row, 2, "Wave 164")
            # A distinct dimension per row keeps (series_id, reference) unique,
            # which the parser requires.
            sheet.cell(row, 3, f"Industry {index}")
            for offset, _ in enumerate(responses):
                sheet.cell(row, 4 + offset, 10.0 + offset)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_absent_question_sheet_is_skipped_not_fatal() -> None:
    """One configured sheet present, the rest rotated out: parse, do not raise."""
    present = min(E.SHEETS)
    body = _workbook([present])
    # The volume floor is what stops the run here, not a missing-sheet error:
    # the synthetic workbook carries headers but no data rows.
    with pytest.raises(ValueError) as caught:
        E.parse_xlsx(body, "digest", "url", datetime.now(UTC))
    assert "question sheet missing" not in str(caught.value)
    assert "unexpectedly short" in str(caught.value)


def test_workbook_with_no_configured_question_fails_loudly() -> None:
    """None of the configured questions present means the layout changed."""
    body = _workbook([])
    with pytest.raises(ValueError, match="none of the configured question sheets"):
        E.parse_xlsx(body, "digest", "url", datetime.now(UTC))


def test_a_single_question_wave_parses_successfully() -> None:
    """The real regression: one rotating question, enough rows, must collect.

    Waves 162 and 164 each carried exactly one configured question. Before the
    per-sheet floor, a wave like this was rejected as "unexpectedly short" even
    though its single question held a complete time series.
    """
    name = min(E.SHEETS)
    observations, catalog, release, _ = E.parse_xlsx(
        _workbook([name], rows=400), "digest", "url", datetime.now(UTC)
    )
    assert len(observations) >= 300
    assert len(catalog) >= 25
    assert release is not None
